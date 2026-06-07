import io
import re
from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from sqlalchemy.orm import Session
from openpyxl import load_workbook
from ..database import get_db
from ..models import Contract, IctService, LegalEntity, Provider
from ..schemas import ImportResult

router = APIRouter()

# Keyword → ICT service category inference
_CATEGORY_KEYWORDS: list[tuple[str, list[str]]] = [
    ('cybersecurity',   ['security', 'endpoint', 'edr', 'av', 'siem']),
    ('data_service',    ['data', 'market', 'feed', 'pricing']),
    ('managed_service', ['telecom', 'internet', 'mobile', 'voip', 'connectivity']),
    ('support',         ['service desk', 'itsm', 'helpdesk']),
    ('IaaS',            ['cloud', 'hosting', 'server', 'infrastructure']),
    ('SaaS',            [
        'portfolio', 'trading', 'investment', 'hr', 'payroll', 'people',
        'accounting', 'finance', 'ledger',
    ]),
]
_CRITICAL_KW = ['portfolio', 'aml', 'compliance', 'screening', 'trading', 'client reporting']


def _infer_category(desc: str) -> str:
    dl = desc.lower()
    for cat, kws in _CATEGORY_KEYWORDS:
        if any(kw in dl for kw in kws):
            return cat
    return 'other'


def _infer_critical(desc: str) -> bool:
    dl = desc.lower()
    return any(kw in dl for kw in _CRITICAL_KW)


def _cell(val) -> str:
    return str(val).strip() if val is not None else ''


def _find_sheet(wb, name: str):
    for sn in wb.sheetnames:
        if sn.strip().lower() == name.strip().lower():
            return wb[sn]
    return None


def _rows(ws):
    it = ws.iter_rows(values_only=True)
    next(it, None)  # skip header
    return it


def _header_index(ws) -> dict[str, int]:
    first = next(ws.iter_rows(values_only=True), [])
    return {_cell(h).strip(): i for i, h in enumerate(first) if h is not None}


@router.post('/excel', response_model=ImportResult)
async def import_excel(file: UploadFile = File(...), db: Session = Depends(get_db)):
    if not (file.filename or '').endswith('.xlsx'):
        raise HTTPException(400, 'Only .xlsx files accepted')

    contents = await file.read()
    wb = load_workbook(io.BytesIO(contents), data_only=True)

    contracts_created = 0
    providers_created = 0
    services_created = 0
    errors: list[str] = []

    # ── Ensure a default LegalEntity exists ──────────────────────────────────
    legal_entity = db.query(LegalEntity).first()
    if not legal_entity:
        legal_entity = LegalEntity(name='Imported Organisation', entity_type='regulated_firm')
        db.add(legal_entity)
        db.flush()

    # ── ROI sheet ─────────────────────────────────────────────────────────────
    roi_ws = _find_sheet(wb, 'ROI')
    if roi_ws:
        for row in _rows(roi_ws):
            if not row:
                continue
            ref   = _cell(row[0]) if len(row) > 0 else ''
            pid   = _cell(row[1]) if len(row) > 1 else ''
            pname = _cell(row[2]) if len(row) > 2 else ''
            desc  = _cell(row[3]) if len(row) > 3 else ''

            if not ref or ref.lower().startswith('contractual'):
                continue
            if db.query(Contract).filter(Contract.contract_ref == ref).first():
                continue  # already imported

            try:
                provider = db.query(Provider).filter(Provider.legal_name == (pname or ref)).first()
                if not provider:
                    provider = Provider(legal_name=pname or ref, lei=pid or None)
                    db.add(provider)
                    db.flush()
                    providers_created += 1
                elif pid and not provider.lei:
                    provider.lei = pid

                contract = Contract(
                    contract_ref=ref,
                    contract_name=(desc[:500] if desc else ref),
                    contract_status='active',
                    provider_id=provider.id,
                    legal_entity_id=legal_entity.id,
                )
                db.add(contract)
                db.flush()
                contracts_created += 1

                service = IctService(
                    service_ref=f'SVC-{ref}',
                    service_description=desc or ref,
                    ict_service_category=_infer_category(desc),
                    contract_id=contract.id,
                    is_critical_or_important=_infer_critical(desc),
                )
                db.add(service)
                services_created += 1

            except Exception as exc:
                errors.append(f'ROI row {ref}: {exc}')

        db.commit()

    # ── Software Asset Register ───────────────────────────────────────────────
    sw_ws = _find_sheet(wb, 'Software Asset Register')
    if sw_ws:
        idx = _header_index(sw_ws)

        def col(row, name: str) -> str:
            i = idx.get(name, -1)
            return _cell(row[i]) if i >= 0 and i < len(row) else ''

        for row in _rows(sw_ws):
            if not row:
                continue
            ref = col(row, 'ROI_Contract_Ref')
            if not ref or ref in ('#N/A', 'None', ''):
                continue
            service = db.query(IctService).filter(IctService.service_ref == f'SVC-{ref}').first()
            if not service:
                continue
            try:
                software = col(row, 'Software')
                used_for = col(row, 'Software Used For')
                classification = col(row, 'Overall_Classification')
                supporting_fn = col(row, 'Supporting Function Classification')
                if software:
                    service.supported_application = software
                if used_for:
                    service.supported_business_process = used_for
                if classification == 'High' and 'Critical or Important' in supporting_fn:
                    service.is_critical_or_important = True
            except Exception as exc:
                errors.append(f'SW row {ref}: {exc}')

        db.commit()

    # ── Dependency Map ────────────────────────────────────────────────────────
    dep_ws = _find_sheet(wb, 'Dependency Map')
    if dep_ws:
        idx = _header_index(dep_ws)

        def col(row, name: str) -> str:
            i = idx.get(name, -1)
            return _cell(row[i]) if i >= 0 and i < len(row) else ''

        for row in _rows(dep_ws):
            if not row:
                continue
            role_id = col(row, 'Role_ID')
            if not re.match(r'^R_\d+$', role_id):
                continue
            role_name = col(row, 'Role Name') or col(row, 'Role_Name')
            refs_raw = col(row, 'ROI_Contract_Ref')
            for ref in [r.strip() for r in refs_raw.split(',') if r.strip()]:
                service = db.query(IctService).filter(
                    IctService.service_ref == f'SVC-{ref}'
                ).first()
                if service and not service.supported_function:
                    service.supported_function = role_name

        db.commit()

    return ImportResult(
        contracts_created=contracts_created,
        providers_created=providers_created,
        services_created=services_created,
        errors=errors,
    )
